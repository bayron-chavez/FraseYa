"""RF-12: importación acotada, sin fórmulas, con rechazo por fila."""
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
import unicodedata
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
import defusedxml  # Exige el parser protegido que openpyxl detecta al importar.


@dataclass
class ResultadoImportacion:
    importadas: int = 0
    rechazadas: list = field(default_factory=list)


def _normalizar(valor):
    return ''.join(c for c in unicodedata.normalize('NFD', str(valor or '').strip().lower())
                   if unicodedata.category(c) != 'Mn')


def importar_excel(ruta, gestion, autenticacion=None, sesion=None):
    ruta = Path(ruta)
    if ruta.suffix.lower() != '.xlsx':
        raise ValueError('Selecciona un archivo .xlsx.')
    with ruta.open('rb') as archivo:
        datos = archivo.read(10 * 1024 * 1024 + 1)
    if len(datos) > 10 * 1024 * 1024:
        raise ValueError('El archivo supera 10 MB.')
    try:
        with ZipFile(BytesIO(datos)) as zip_excel:
            archivos = zip_excel.infolist()
            if (len(archivos) > 1000 or sum(a.file_size for a in archivos) > 50 * 1024 * 1024
                    or any(a.flag_bits & 1 or 'vbaproject' in a.filename.lower() for a in archivos)):
                raise ValueError('El archivo contiene elementos no permitidos o es demasiado grande.')
        libro = load_workbook(BytesIO(datos), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError) as error:
        raise ValueError('El archivo Excel no es válido.') from error
    resultado = ResultadoImportacion()
    try:
        hoja = libro.active
        if hoja.max_row and hoja.max_row > 10001:
            raise ValueError('El archivo supera el límite de 10000 filas de datos.')
        filas = hoja.iter_rows(max_row=10002, max_col=4)
        cabecera = next(filas)
        if [_normalizar(c.value) for c in cabecera] != ['categoria', 'titulo', 'abreviatura', 'contenido']:
            raise ValueError('Las columnas deben ser: categoría, título, abreviatura, contenido, en ese orden.')
        categorias = {c['nombre'].casefold(): c for c in gestion.categorias_propias()}
        abreviaturas = {f['abreviatura'].casefold() for f in gestion.repo.listar_frases()}
        admin_verificado = False
        for numero, fila in enumerate(filas, start=2):
            if all(c.value is None for c in fila):
                continue
            if numero > 10001:
                raise ValueError('El archivo supera el límite de filas.')
            try:
                if any(c.data_type == 'f' or not isinstance(c.value, str) for c in fila):
                    raise ValueError('Usa texto en las cuatro columnas; no se admiten fórmulas.')
                categoria, titulo, abreviatura, contenido = (c.value.strip() for c in fila)
                if not categoria or len(categoria) > 200:
                    raise ValueError('Categoría vacía o demasiado larga.')
                if not abreviatura or any(c.isspace() for c in abreviatura):
                    raise ValueError('La abreviatura no puede estar vacía ni contener espacios.')
                if abreviatura.casefold() in abreviaturas:
                    raise ValueError('La abreviatura ya está en uso.')
                if not titulo or len(titulo) > 500 or not contenido or len(contenido) > 100000 or len(abreviatura) > 100:
                    raise ValueError('Título, abreviatura o contenido vacío o demasiado largo.')
                cat = categorias.get(categoria.casefold())
                if cat is None:
                    if autenticacion is None or sesion is None:
                        raise PermissionError('La categoría no existe; solo el administrador puede crearla.')
                    if not admin_verificado:
                        autenticacion.validar(sesion, administrador=True)
                        admin_verificado = True
                with gestion.repo.db:
                    if cat is None:
                        ident = gestion.repo.db.execute('INSERT INTO CATEGORIA(catalogo_id,nombre,color) VALUES(?,?,?)',
                            (gestion._catalogo_propio(), categoria, '#64748B')).lastrowid
                        cat = {'id': ident, 'nombre': categoria}
                    gestion.repo._crear_frase(cat['id'], titulo, abreviatura, contenido, 'propia')
                categorias[categoria.casefold()] = cat
                abreviaturas.add(abreviatura.casefold())
                resultado.importadas += 1
            except (ValueError, PermissionError) as error:
                resultado.rechazadas.append((numero, str(error)))
    finally:
        libro.close()
    return resultado
